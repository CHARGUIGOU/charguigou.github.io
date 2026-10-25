#!/usr/bin/env python3
"""
Génère _data/characters.yml (la base des personnages) à partir d'UN tableau Excel :
    _data/characters.xlsx
        feuille « Personnages » : une ligne = un personnage
        feuille « Apparitions » : une ligne = un film où le personnage apparaît

Le fichier characters.yml est RECRÉÉ à chaque build (étape du workflow GitHub, voir LISEZMOI.md).
Les réglages communs (image_base : début des URLs d'images) restent lus dans le characters.yml existant.

Usage (depuis la racine du dépôt) :
    python _scripts/generate_characters.py            # écrit _data/characters.yml
    python _scripts/generate_characters.py --check    # vérifie le tableau, n'écrit rien
"""
import argparse, datetime, os, re, sys

try:
    import yaml
except ImportError:
    sys.exit("PyYAML manquant : pip install pyyaml")
try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl manquant : pip install openpyxl")

DEFAULT_BASE = "https://res.cloudinary.com/dofvjifeo/image/upload/c_limit,w_320,q_auto,f_auto/"
ID_OK = re.compile(r"^[a-z0-9][a-z0-9_-]*$")
YES = {"yes", "y", "oui", "o", "true", "vrai", "x", "1"}
NO = {"no", "n", "non", "false", "faux", "0", "hidden", "masqué", "masque"}


def norm_header(h):
    return re.sub(r"[\s\-]+", "_", str(h or "").strip().lower())


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return str(v.year)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def split_list(t):
    return [x.strip() for x in re.split(r"[;\r\n]+", t) if x.strip()]


def read_sheet(wb, names):
    for n in names:
        if n in wb.sheetnames:
            ws = wb[n]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return [], []
            heads = [norm_header(h) for h in rows[0]]
            return heads, rows[1:]
    return None, None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=".")
    ap.add_argument("--table", default=None)
    ap.add_argument("--out", default=None)
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    root = os.path.abspath(a.root)
    table = a.table or os.path.join(root, "_data", "characters.xlsx")
    out_path = a.out or os.path.join(root, "_data", "characters.yml")
    if not os.path.exists(table):
        print("Aucun tableau _data/characters.xlsx : characters.yml reste tel quel.")
        return

    wb = openpyxl.load_workbook(table, read_only=True, data_only=True)
    ph, prows = read_sheet(wb, ("Personnages", "Characters"))
    ah, arows = read_sheet(wb, ("Apparitions", "Appearances"))
    if ph is None or ah is None:
        sys.exit("ERREUR : il faut deux feuilles nommées « Personnages » et « Apparitions ».")
    for need, heads, label in (("id", ph, "Personnages"), ("character", ah, "Apparitions")):
        if need not in heads:
            sys.exit("ERREUR : la colonne « %s » est introuvable dans la feuille « %s »." % (need, label))

    problems = []
    chars, order = {}, []
    for n, r in enumerate(prows, start=2):
        row = {h: cell_text(r[i]) if i < len(r) else "" for i, h in enumerate(ph) if h and not h.startswith(("#", "_"))}
        cid = row.get("id", "")
        if not any(row.values()) or cid.startswith("#"):
            continue
        if not cid:
            problems.append("Personnages, ligne %d : identifiant (id) vide, ligne ignorée" % n)
            continue
        if not ID_OK.match(cid):
            problems.append("Personnages, ligne %d : identifiant « %s » invalide (minuscules, chiffres, - et _, sans accent ni espace), ligne ignorée" % (n, cid))
            continue
        if cid in chars:
            problems.append("Personnages, ligne %d : identifiant « %s » en double, ligne ignorée" % (n, cid))
            continue
        item = {"id": cid}
        if row.get("name"):
            item["name"] = row["name"]
        else:
            problems.append("Personnages, ligne %d (%s) : nom (name) vide" % (n, cid))
        for k in ("img", "franchise"):
            if row.get(k):
                item[k] = row[k]
        if row.get("keywords"):
            item["keywords"] = split_list(row["keywords"])
        if row.get("link"):
            item["link"] = row["link"]
        vis = row.get("visible", "").lower()
        if vis in NO:
            item["hidden"] = True
        elif vis and vis not in YES:
            problems.append("Personnages, ligne %d (%s) : visible « %s » non compris (écris yes ou no), personnage laissé visible" % (n, cid, row["visible"]))
        item["appearances"] = []
        chars[cid] = item
        order.append(cid)

    n_app = 0
    for n, r in enumerate(arows, start=2):
        row = {h: cell_text(r[i]) if i < len(r) else "" for i, h in enumerate(ah) if h and not h.startswith(("#", "_"))}
        cid = row.get("character", "")
        if not any(row.values()) or cid.startswith("#"):
            continue
        if cid not in chars:
            problems.append("Apparitions, ligne %d : personnage « %s » introuvable dans la feuille Personnages, ligne ignorée" % (n, cid))
            continue
        if not row.get("film"):
            problems.append("Apparitions, ligne %d (%s) : film vide, ligne ignorée" % (n, cid))
            continue
        ap_item = {"film": row["film"]}
        ap_item["year"] = row.get("year", "")
        if not ap_item["year"]:
            problems.append("Apparitions, ligne %d (%s / %s) : année vide" % (n, cid, row["film"]))
        if row.get("keywords"):
            ap_item["keywords"] = split_list(row["keywords"])
        if row.get("img"):
            ap_item["img"] = row["img"]
        lead = row.get("lead", "")
        if lead:
            if re.fullmatch(r"\d+", lead) and lead not in ("0",):
                ap_item["lead"] = int(lead)
            elif lead.lower() in YES:
                ap_item["lead"] = True
            elif lead.lower() not in NO:
                problems.append("Apparitions, ligne %d (%s / %s) : lead « %s » non compris (yes, ou un nombre : 1 = en premier), ignoré" % (n, cid, row["film"], lead))
        if row.get("upcoming", "").lower() in YES:
            ap_item["upcoming"] = True
        chars[cid]["appearances"].append(ap_item)
        n_app += 1

    base, extra = DEFAULT_BASE, {}
    if os.path.exists(out_path):
        try:
            old = yaml.safe_load(open(out_path, encoding="utf-8-sig").read())
            if isinstance(old, dict):
                base = old.get("image_base", base)
                extra = {k: v for k, v in old.items() if k not in ("image_base", "list")}
        except yaml.YAMLError:
            pass

    hidden = sum(1 for c in chars.values() if c.get("hidden"))
    print("tableau : %s" % os.path.relpath(table, root))
    print("personnages : %d (dont %d masqués) | apparitions : %d | problèmes : %d" % (len(order), hidden, n_app, len(problems)))
    for p in problems:
        print("  -", p)
    if a.check:
        return
    data = {"image_base": base}
    data.update(extra)
    data["list"] = [chars[c] for c in order]
    head = ("# Fichier GÉNÉRÉ depuis _data/characters.xlsx à chaque build : ne le modifie pas à la main,\n"
            "# modifie le tableau (feuilles Personnages et Apparitions).\n")
    txt = head + yaml.safe_dump(data, allow_unicode=True, sort_keys=False, default_flow_style=False, width=10000)
    with open(out_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(txt)
    print("écrit : %s" % os.path.relpath(out_path, root))


if __name__ == "__main__":
    main()
